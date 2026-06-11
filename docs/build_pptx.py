# -*- coding: utf-8 -*-
"""GoGlobal — PPTX presentation builder (expanded, 36 slides, Cyrillic-safe)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- palette (matches HTML deck & CRM v5 design system) ----------
BG      = RGBColor(0x0B, 0x0C, 0x0F)
BG2     = RGBColor(0x12, 0x14, 0x19)
INK     = RGBColor(0xE7, 0xE9, 0xEE)
INK2    = RGBColor(0x9A, 0xA1, 0xAE)
INK3    = RGBColor(0x6E, 0x75, 0x82)
LINE    = RGBColor(0x23, 0x26, 0x2E)
ACCENT  = RGBColor(0x63, 0x66, 0xF1)
ACCENT2 = RGBColor(0x81, 0x8C, 0xF8)
OK      = RGBColor(0x10, 0xB9, 0x81)
OKL     = RGBColor(0x86, 0xEF, 0xAC)
WARN    = RGBColor(0xF5, 0x9E, 0x0B)
WARNL   = RGBColor(0xFD, 0xE6, 0x8A)
BAD     = RGBColor(0xEF, 0x44, 0x44)
BADL    = RGBColor(0xFD, 0xA4, 0xAF)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SKYL    = RGBColor(0xA5, 0xB4, 0xFC)

FONT = "Arial"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height
TOTAL = 36
_n = [0]

# ---------- helpers ----------
def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    # top accent strip
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Pt(4))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background(); strip.shadow.inherit = False
    _n[0] += 1
    # page number
    tb = s.shapes.add_textbox(SLIDE_W - Inches(1.4), Inches(0.18), Inches(1.2), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run(); run.text = f"{_n[0]:02d} / {TOTAL}"
    run.font.size = Pt(10); run.font.color.rgb = INK3; run.font.name = FONT
    return s

def txt(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """lines: list of (text, size, color, bold) or list of lists of runs [(text,size,color,bold),...]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        runs = line if isinstance(line, list) else [line]
        for (t, size, color, bold) in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = FONT
    return tb

def kicker(slide, text, x=Inches(0.7), y=Inches(0.42)):
    txt(slide, x, y, Inches(9), Inches(0.32),
        [(text.upper(), 11, ACCENT, True)])

def title(slide, text, x=Inches(0.7), y=Inches(0.72), size=30, color=INK, w=None):
    txt(slide, x, y, w or (SLIDE_W - x - Inches(0.7)), Inches(0.7),
        [(text, size, color, True)])

def card(slide, x, y, w, h, fill=BG2, line_c=LINE):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.06
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line_c; c.line.width = Pt(0.75)
    c.shadow.inherit = False
    return c

def footer(slide, right="GoGlobal"):
    txt(slide, Inches(0.7), SLIDE_H - Inches(0.42), Inches(6), Inches(0.3),
        [("GoGlobal · база знаний", 9, INK3, False)])
    txt(slide, SLIDE_W - Inches(5.7), SLIDE_H - Inches(0.42), Inches(5), Inches(0.3),
        [(right, 9, INK3, False)], align=PP_ALIGN.RIGHT)

def style_cell(cell, text_runs, fill=None, align=PP_ALIGN.LEFT, size=11):
    """text_runs: str or list of (text,color,bold)"""
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = BG2
    tf = cell.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(6); tf.margin_right = Pt(6)
    tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
    p = tf.paragraphs[0]; p.alignment = align
    if isinstance(text_runs, str):
        text_runs = [(text_runs, INK, False)]
    for (t, color, bold) in text_runs:
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.name = FONT

def make_table(slide, x, y, w, rows_data, col_widths, header=True, size=11, row_h=Pt(24)):
    """rows_data: list of rows; each row list of cells; cell = str | (str,color) | (str,color,bold) | list of runs"""
    nrows = len(rows_data); ncols = len(rows_data[0])
    gfx = slide.shapes.add_table(nrows, ncols, x, y, w, row_h * nrows)
    table = gfx.table
    # disable default style banding
    tbl = gfx._element.graphic.graphicData.tbl
    for el in tbl.findall(qn('a:tblPr')):
        el.set('bandRow', '0'); el.set('firstRow', '0')
    total = sum(col_widths)
    for i, cw in enumerate(col_widths):
        table.columns[i].width = int(w * cw / total)
    for ri, row in enumerate(rows_data):
        table.rows[ri].height = row_h
        for ci, cellv in enumerate(row):
            cell = table.cell(ri, ci)
            if header and ri == 0:
                runs = [(str(cellv), INK2, True)]
                style_cell(cell, runs, fill=BG2, size=size-1)
                continue
            if isinstance(cellv, list):
                runs = cellv
            elif isinstance(cellv, tuple):
                t = cellv[0]; color = cellv[1] if len(cellv) > 1 else INK
                bold = cellv[2] if len(cellv) > 2 else False
                runs = [(t, color, bold)]
            else:
                runs = [(str(cellv), INK, False)]
            fill = BG if ri % 2 == 1 else RGBColor(0x0D, 0x16, 0x28)
            style_cell(cell, runs, fill=fill, size=size)
    return table

def kpi_box(slide, x, y, w, h, label, value, sub="", val_color=INK):
    card(slide, x, y, w, h)
    lines = [(label.upper(), 9, INK3, True), (value, 22, val_color, True)]
    if sub: lines.append((sub, 9, INK3, False))
    txt(slide, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2), lines)

def bullet_list(slide, x, y, w, h, items, size=12, gap_para=True, color=INK2, bcolor=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if gap_para: p.space_after = Pt(4)
        r = p.add_run(); r.text = "•  "
        r.font.size = Pt(size); r.font.color.rgb = bcolor or ACCENT; r.font.bold = True; r.font.name = FONT
        if isinstance(it, list):
            for (t, c, b) in it:
                r2 = p.add_run(); r2.text = t
                r2.font.size = Pt(size); r2.font.color.rgb = c; r2.font.bold = b; r2.font.name = FONT
        else:
            r2 = p.add_run(); r2.text = it
            r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.bold = False; r2.font.name = FONT
    return tb

def steps_list(slide, x, y, w, items, size=12, gap=Inches(0.62)):
    for i, (head, body) in enumerate(items):
        yy = y + gap * i
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, yy, Inches(0.32), Inches(0.32))
        circ.fill.solid(); circ.fill.fore_color.rgb = BG2
        circ.line.color.rgb = ACCENT; circ.line.width = Pt(1)
        circ.shadow.inherit = False
        tfc = circ.text_frame; pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
        rc = pc.add_run(); rc.text = str(i+1)
        rc.font.size = Pt(11); rc.font.bold = True; rc.font.color.rgb = ACCENT; rc.font.name = FONT
        txt(slide, x + Inches(0.48), yy - Inches(0.04), w - Inches(0.5), gap,
            [(head, size, INK, True), (body, size-1.5, INK2, False)])

# ============================================================
# 1 · COVER
# ============================================================
s = add_slide()
# gradient-ish accents
for (gx, gy, gw, gh, col) in [(0, 0, Inches(5), Inches(3), RGBColor(0x14, 0x2A, 0x52)),
                              (SLIDE_W - Inches(4), SLIDE_H - Inches(2.5), Inches(4), Inches(2.5), RGBColor(0x0C, 0x31, 0x40))]:
    g = s.shapes.add_shape(MSO_SHAPE.OVAL, gx, gy, gw, gh)
    g.fill.solid(); g.fill.fore_color.rgb = col
    g.line.fill.background(); g.shadow.inherit = False

logo = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(1.0), Inches(1.0))
logo.adjustments[0] = 0.25
logo.fill.solid(); logo.fill.fore_color.rgb = WHITE
logo.line.fill.background(); logo.shadow.inherit = False
pl = logo.text_frame.paragraphs[0]; pl.alignment = PP_ALIGN.CENTER
rl = pl.add_run(); rl.text = "GG"; rl.font.size = Pt(28); rl.font.bold = True
rl.font.color.rgb = BG; rl.font.name = FONT
logo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

txt(s, Inches(0.8), Inches(2.2), Inches(11), Inches(1.2), [("GoGlobal", 60, WHITE, True)])
txt(s, Inches(0.8), Inches(3.35), Inches(11), Inches(1.0),
    [("Платформа агентства образования за рубежом:", 20, INK2, False),
     ("сайт + админ-панель + CRM в одной системе", 20, INK2, False)])
meta = [("Адрес", "goglobal.kg"), ("Для кого", "Админ · РОП · Продавцы"),
        ("Версия", "4.1 · июнь 2026"), ("Слайдов", "36 · ~25 минут")]
for i, (k, v) in enumerate(meta):
    x = Inches(0.8) + Inches(3.05) * i
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(5.4), Pt(2.5), Inches(0.75))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background(); bar.shadow.inherit = False
    txt(s, x + Inches(0.15), Inches(5.4), Inches(2.8), Inches(0.8),
        [(k.upper(), 10, INK3, True), (v, 14, INK, True)])

# ============================================================
# 2 · TOC
# ============================================================
s = add_slide()
kicker(s, "Содержание")
title(s, "О чём эта презентация")
toc_l = [("I · Платформа в цифрах", "03"), ("II · Три части системы", "04–05"),
         ("III · Как всё связано", "06–07"), ("IV · Роли и права", "08–09"),
         ("V · Сайт по экранам", "10–13"), ("VI · CRM продавца", "14–20")]
toc_r = [("VII · Админка и аналитика", "21–27"), ("VIII · Путь клиента", "28–29"),
         ("IX · Истории из практики", "30–32"), ("X · Карта действий", "33"),
         ("XI · Чек-листы", "34–35"), ("Финал", "36")]
for col, items, x in ((0, toc_l, Inches(0.7)), (1, toc_r, Inches(6.9))):
    for i, (name, pg) in enumerate(items):
        y = Inches(1.7) + Inches(0.52) * i
        txt(s, x, y, Inches(4.7), Inches(0.45), [(name, 14, INK, True)])
        txt(s, x + Inches(4.8), y, Inches(0.9), Inches(0.45), [(pg, 12, INK3, False)], align=PP_ALIGN.RIGHT)
card(s, Inches(0.7), Inches(5.2), Inches(11.9), Inches(1.3))
txt(s, Inches(0.95), Inches(5.4), Inches(11.4), Inches(1.0),
    [("КАК ЧИТАТЬ", 10, ACCENT, True),
     ("Продавцам достаточно разделов V–VI и IX. РОПу — VII–XI. Админу — всё. Каждый слайд самодостаточен.", 12, INK2, False)])
footer(s, "Содержание")

# ============================================================
# 3 · NUMBERS
# ============================================================
s = add_slide()
kicker(s, "I · Платформа в цифрах")
title(s, "Что построено")
kpis = [("Частей системы", "3", "сайт · админка · CRM"),
        ("Ролей", "4", "админ · РОП · менеджер · партнёр"),
        ("Разделов админки", "30+", "аналитика · операции · контент"),
        ("Виджетов аналитики", "11", "без единого Excel"),
        ("Видов списка лидов", "5", "карточки · таблица · 2 канбана · календарь"),
        ("Вкладок в карточке", "7", "обзор · сделка · задачи · файлы · чат · аудит · связи"),
        ("Статусов воронки", "7+9", "обработка + этапы клиента"),
        ("Авто-защит", "6", "SLA · дубли · сценарии · отчёты · аудит · бэкап")]
for i, (l, v, sub) in enumerate(kpis):
    x = Inches(0.7) + Inches(3.05) * (i % 4)
    y = Inches(1.55) + Inches(1.15) * (i // 4)
    kpi_box(s, x, y, Inches(2.85), Inches(1.0), l, v, sub)
rows = [["Было (без платформы)", "Стало (с платформой)"],
        ["CRM-аренда — платёж каждый месяц", ("Своя CRM — 0 платежей за лицензии", OKL)],
        ["Заявки на почту, теряются", ("Заявка у менеджера за секунды + push", OKL)],
        ["Отчёты руками в Excel по пятницам", ("Цифры считаются сами, в реальном времени", OKL)],
        ["«Кто вёл клиента?» — никто не помнит", ("Полная история: статусы, звонки, файлы, аудит", OKL)],
        ["Контент сайта меняет программист", ("Меняет админ сам за 5–10 минут", OKL)]]
make_table(s, Inches(0.7), Inches(4.05), Inches(11.9), rows, [1, 1.2], size=11, row_h=Pt(22))
footer(s, "I · Платформа в цифрах")

# ============================================================
# 4 · THREE PARTS
# ============================================================
s = add_slide()
kicker(s, "II · Три части системы")
title(s, "Один движок — три «двери»")
parts = [("Сайт", "goglobal.kg", "Маркетинговая витрина: направления, вузы, отзывы, калькулятор, форма заявки. Открыт всем.", "Для клиентов", ACCENT),
         ("Админ-панель", "goglobal.kg/admin", "Сотрудники, воронка, шаблоны, авто-правила, контент сайта. Вся аналитика — 11 виджетов.", "Для админа и РОПа", RGBColor(0x63, 0x66, 0xF1)),
         ("CRM", "goglobal.kg/lidy", "Рабочее место менеджера: лиды, карточки, задачи, документы, календарь, шаблоны WhatsApp.", "Для продавцов", OK)]
for i, (name, url, desc, who, color) in enumerate(parts):
    x = Inches(0.7) + Inches(4.0) * i
    card(s, x, Inches(1.6), Inches(3.8), Inches(2.25), line_c=color)
    txt(s, x + Inches(0.2), Inches(1.75), Inches(3.4), Inches(2.0),
        [(who.upper(), 9, color, True), (name, 18, INK, True), (url, 11, INK3, False), (desc, 11, INK2, False)])
rows = [["Параметр", "Сайт", "Админка", "CRM"],
        ["Нужен пароль", "Нет", "Да (пароль админа)", "Да (личный логин)"],
        ["Кто пользуется", "Клиенты", "Админ, РОП", "Менеджеры, РОП, партнёры"],
        ["Главное действие", "Оставить заявку", "Настроить и проанализировать", "Превратить лид в контракт"],
        ["Как часто заходят", "1 раз (клиент)", "1–2 раза в день", "Весь рабочий день"]]
make_table(s, Inches(0.7), Inches(4.15), Inches(11.9), rows, [0.9, 0.8, 1.1, 1.2], size=11, row_h=Pt(24))
footer(s, "II · Три части")

# ============================================================
# 5 · LOOK & FEEL
# ============================================================
s = add_slide()
kicker(s, "II · Три части системы")
title(s, "Как они выглядят")
looks = [("Сайт", "Яркий, продающий: синие градиенты, оранжевые акценты, крупные заголовки, фото студентов. Дизайн для эмоций клиента.", ACCENT),
         ("Админка", "Спокойная и минималистичная: тёмный фон, разделы-аккордеоны, ничего лишнего. Дизайн для долгой работы без усталости глаз.", RGBColor(0x63, 0x66, 0xF1)),
         ("CRM", "Рабочий стол: виджет «Мой день» сверху, список лидов с цветными статусами, всё в 1–2 клика. Дизайн для скорости.", OK)]
for i, (name, desc, color) in enumerate(looks):
    x = Inches(0.7) + Inches(4.0) * i
    card(s, x, Inches(1.6), Inches(3.8), Inches(2.6), line_c=color)
    txt(s, x + Inches(0.2), Inches(1.8), Inches(3.4), Inches(2.3),
        [(name, 17, INK, True), (desc, 11.5, INK2, False)])
card(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(1.7))
txt(s, Inches(0.95), Inches(4.7), Inches(11.4), Inches(1.4),
    [("ДИЗАЙН-СИСТЕМА РАБОЧИХ ЭКРАНОВ (V5)", 10, ACCENT, True),
     ("Нейтральный графитовый фон + один спокойный акцент (индиго) + вторичный теал. Красный — только тревога, зелёный — успех.", 12, INK2, False),
     ("Статусы — выпадающий список с подсказками, этапы — степпер 1→9, у ключевых цифр значок ⓘ с пояснением. Без неона и свечений.", 12, INK2, False)])
footer(s, "II · Как выглядят")

# ============================================================
# 6 · LEAD FLOW
# ============================================================
s = add_slide()
kicker(s, "III · Как всё связано")
title(s, "Путь заявки за 60 секунд")
flow = [("🌍  Сайт", "клиент заполнил форму"), ("📥  CRM", "лид у менеджера + push"), ("📊  Админка", "виден в аналитике РОПа")]
for i, (name, desc) in enumerate(flow):
    x = Inches(0.9) + Inches(4.2) * i
    card(s, x, Inches(1.55), Inches(3.3), Inches(0.95))
    txt(s, x + Inches(0.2), Inches(1.67), Inches(3.0), Inches(0.8),
        [(name, 15, INK, True), (desc, 11, INK3, False)])
    if i < 2:
        txt(s, x + Inches(3.4), Inches(1.7), Inches(0.7), Inches(0.6), [("→", 26, ACCENT, True)], align=PP_ALIGN.CENTER)
rows = [["Сек", "Событие", "Кто видит"],
        ["0", "Клиент нажал «Отправить заявку»", "Клиент: «Спасибо! Менеджер свяжется»"],
        ["~1", "Проверка на дубликат (тот же телефон/email?)", "Дубль → комментарий к старому лиду"],
        ["~1", "Подбор менеджера: правила, потом очередь", "—"],
        ["~2", "Лид в CRM со статусом «Новый»", "Менеджер: push-уведомление"],
        ["~2", "Запустился таймер SLA (3 часа на ответ)", "Таймер в карточке лида"],
        ["~3", "Сообщение в Telegram-чат отдела", "РОП и команда"]]
make_table(s, Inches(0.7), Inches(2.75), Inches(11.9), rows, [0.35, 1.6, 1.3], size=11, row_h=Pt(23))
card(s, Inches(0.7), Inches(5.65), Inches(5.85), Inches(1.25))
txt(s, Inches(0.95), Inches(5.8), Inches(5.4), Inches(1.0),
    [("НЕ ОТВЕТИЛ ВОВРЕМЯ?", 9, WARNL, True),
     ("Через 3 ч: лид краснеет, РОПу сигнал в Telegram, минус в личную статистику менеджера.", 11, INK2, False)])
card(s, Inches(6.75), Inches(5.65), Inches(5.85), Inches(1.25))
txt(s, Inches(7.0), Inches(5.8), Inches(5.4), Inches(1.0),
    [("ВСЕ ОФЛАЙН?", 9, WARNL, True),
     ("Лид встаёт в очередь. Кто-то включил «в сети» — очередь раздаётся автоматически.", 11, INK2, False)])
footer(s, "III · Путь заявки")

# ============================================================
# 7 · WHAT CONTROLS WHAT
# ============================================================
s = add_slide()
kicker(s, "III · Как всё связано")
title(s, "Что чем управляется")
rows1 = [["На сайте", "Управляется в админке"],
         ["Главное фото и тексты", "«Главный экран (Hero)»"],
         ["Страны и вузы", "«Страны и Университеты»"],
         ["Цены калькулятора", "«Калькулятор стоимости обучения»"],
         ["Отзывы студентов", "«Отзывы»"],
         ["Телефон, адрес, график", "«Контактная информация»"],
         ["Секции вкл/выкл", "«Видимость блоков»"]]
rows2 = [["В CRM", "Управляется в админке"],
         ["Кто может войти", "«Менеджеры»"],
         ["Статусы воронки", "«Статусы лидов»"],
         ["Шаблоны WhatsApp", "«Шаблоны быстрых ответов»"],
         ["Кому падают лиды", "«Авто-распределение»"],
         ["Реакции на бездействие", "«Авто-сценарии»"],
         ["Статьи для менеджеров", "«База знаний»"]]
txt(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(0.4), [("Сайт ← Админка", 16, INK, True)])
make_table(s, Inches(0.7), Inches(1.95), Inches(5.8), rows1, [1, 1.1], size=10.5, row_h=Pt(23))
txt(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.4), [("CRM ← Админка", 16, INK, True)])
make_table(s, Inches(6.8), Inches(1.95), Inches(5.8), rows2, [1, 1.1], size=10.5, row_h=Pt(23))
card(s, Inches(0.7), Inches(5.65), Inches(11.9), Inches(1.15))
txt(s, Inches(0.95), Inches(5.82), Inches(11.4), Inches(0.9),
    [("ПРИНЦИП", 10, ACCENT, True),
     ("Менеджеры работают, админ настраивает, РОП анализирует. Продавец не может сломать настройки, админ не вмешивается в сделки.", 12, INK2, False)])
footer(s, "III · Управление")

# ============================================================
# 8 · ROLES
# ============================================================
s = add_slide()
kicker(s, "IV · Роли и права")
title(s, "Четыре роли")
roles = [("👑  Администратор", "Хозяин системы · 1-2 человека",
          ["Создаёт сотрудников, назначает роли", "Меняет контент сайта без программиста",
           "Настраивает воронку, шаблоны, авто-правила", "Резервные копии раз в месяц"]),
         ("📈  РОП / Тимлид", "Руководитель отдела продаж",
          ["Видит лиды всех менеджеров", "Переназначает, передаёт, удаляет лиды",
           "Утренний отчёт в Telegram в 09:00", "1-на-1 встречи по готовым отчётам"]),
         ("💼  Менеджер", "Продавец · основная роль",
          ["Видит только свои лиды", "Чат, задачи, документы, статусы, сделки",
           "Шаблоны WhatsApp в один клик", "Передача лида коллеге (10 мин на принятие)"]),
         ("🤝  Партнёр", "Внешний агент · рефералы",
          ["Только создаёт лидов", "Видит только тех, кого привёл сам",
           "Отслеживает свою комиссию", "Нет доступа к чужой базе"])]
for i, (name, sub, items) in enumerate(roles):
    x = Inches(0.7) + Inches(6.1) * (i % 2)
    y = Inches(1.55) + Inches(2.6) * (i // 2)
    card(s, x, y, Inches(5.9), Inches(2.4))
    txt(s, x + Inches(0.25), y + Inches(0.15), Inches(5.4), Inches(0.7),
        [(name, 16, INK, True), (sub, 10.5, INK3, False)])
    bullet_list(s, x + Inches(0.25), y + Inches(0.85), Inches(5.4), Inches(1.5), items, size=10.5)
footer(s, "IV · Роли")

# ============================================================
# 9 · RIGHTS MATRIX
# ============================================================
s = add_slide()
kicker(s, "IV · Роли и права")
title(s, "Матрица прав — кто что может")
Y = lambda: ("✓", OKL, True); N = lambda: ("—", INK3, False)
rows = [["Действие", "Админ", "РОП", "Менеджер", "Партнёр"],
        ["Вход в админ-панель", Y(), Y(), N(), N()],
        ["Вход в CRM", N(), Y(), Y(), Y()],
        ["Видеть все лиды", ("через админку", WARNL), Y(), ("только свои", INK3), ("только свои", INK3)],
        ["Создавать лидов вручную", N(), Y(), Y(), Y()],
        ["Менять статус лида", N(), ("✓ любого", OKL, True), ("✓ своего", OKL, True), N()],
        ["Переназначать лиды", N(), Y(), ("передача с согласием", WARNL), N()],
        ["Удалять лиды", N(), Y(), N(), N()],
        ["Создавать сотрудников", Y(), N(), N(), N()],
        ["Менять контент сайта", Y(), N(), N(), N()],
        ["Воронка и шаблоны", Y(), N(), N(), N()],
        ["Смотреть аналитику", ("✓ вся", OKL, True), ("✓ вся", OKL, True), ("свой «Мой день»", WARNL), ("своя комиссия", WARNL)],
        ["Резервные копии", Y(), N(), N(), N()]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [1.5, 0.8, 0.8, 1.0, 0.9], size=10.5, row_h=Pt(24))
footer(s, "IV · Матрица прав")

# ============================================================
# 10 · SITE SECTIONS
# ============================================================
s = add_slide()
kicker(s, "V · Сайт по экранам")
title(s, "Семь секций сверху вниз")
rows = [["#", "Секция", "Что в ней", "Зачем"],
        ["1", ("Главный экран", INK, True), "Заголовок, бейдж, 2 кнопки, фото", "Зацепить за 5 секунд"],
        ["2", ("О компании", INK, True), "Миссия, цифры (10+ лет, 500+ вузов), 4 услуги", "Доверие"],
        ["3", ("Партнёры", INK, True), "Тёмный блок с логотипами вузов", "Статус и престиж"],
        ["4", ("Направления", INK, True), "Страны с фото-каруселями вузов", "Выбор страны мечты"],
        ["5", ("Калькулятор", INK, True), "Страна → ориентир стоимости/год", "Прозрачность цен"],
        ["6", ("Отзывы", INK, True), "Бегущая лента историй студентов", "Социальное доказательство"],
        ["7", ("FAQ + Заявка", INK, True), "Вопросы-ответы + форма с контактами", "Снять сомнения → лид"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [0.25, 1.0, 2.0, 1.3], size=11, row_h=Pt(25))
card(s, Inches(0.7), Inches(4.95), Inches(5.85), Inches(1.9))
txt(s, Inches(0.95), Inches(5.1), Inches(5.4), Inches(0.35), [("КОНВЕРСИОННЫЕ ТОЧКИ", 10, ACCENT, True)])
bullet_list(s, Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.3),
            ["Кнопка «Заявка» в шапке — всегда на виду", "«Выбрать ВУЗ» на главном экране",
             "«Подробнее» у каждой страны", "Плавающая кнопка WhatsApp — всегда"], size=10.5)
card(s, Inches(6.75), Inches(4.95), Inches(5.85), Inches(1.9))
txt(s, Inches(7.0), Inches(5.1), Inches(5.4), Inches(1.6),
    [("МОБИЛЬНАЯ ВЕРСИЯ", 10, ACCENT, True),
     ("Сайт полностью адаптивен: «бургер»-меню, карточки в столбик, кнопки на всю ширину.", 11, INK2, False),
     ("Большинство клиентов приходят с телефона — это учтено.", 11, INK2, False)])
footer(s, "V · Сайт")

# ============================================================
# 11 · CALCULATOR
# ============================================================
s = add_slide()
kicker(s, "V · Сайт по экранам")
title(s, "Калькулятор стоимости")
txt(s, Inches(0.7), Inches(1.45), Inches(6.0), Inches(0.8),
    [("Клиент выбирает страну — мгновенно видит ориентир «от $X в год».", 14, INK, False),
     ("Проверено в живом тесте (июнь 2026):", 12, INK3, False)])
rows = [["Страна", "Ориентир / год"],
        ["США", ("$35 000", INK, True)],
        ["Япония", ("$15 000", INK, True)],
        ["Китай, Германия, Испания…", ("свои цифры", INK3)]]
make_table(s, Inches(0.7), Inches(2.45), Inches(5.8), rows, [1.2, 1], size=12, row_h=Pt(27))
card(s, Inches(0.7), Inches(4.55), Inches(5.8), Inches(2.1))
txt(s, Inches(0.95), Inches(4.7), Inches(5.3), Inches(0.35), [("ИЗ ЧЕГО СКЛАДЫВАЕТСЯ СУММА", 10, ACCENT, True)])
bullet_list(s, Inches(0.95), Inches(5.05), Inches(5.3), Inches(1.5),
            [[("Обучение", INK, True), (" — самый доступный вуз страны", INK2, False)],
             [("Проживание", INK, True), (" — минимальная оценка за год", INK2, False)],
             [("Услуги агентства", INK, True), (" — если настроены", INK2, False)],
             [("Галочка «грант»", INK, True), (" — обучение вычитается", INK2, False)]], size=11)
# mock result card
mock = card(s, Inches(7.3), Inches(1.9), Inches(5.0), Inches(3.0), fill=RGBColor(0x10, 0x2A, 0x52), line_c=RGBColor(0x1E, 0x3A, 0x8A))
txt(s, Inches(7.3), Inches(2.4), Inches(5.0), Inches(2.2),
    [("ОРИЕНТИРОВОЧНО ОТ", 11, RGBColor(0x93, 0xC5, 0xFD), True),
     ("$15 000", 48, WHITE, True),
     ("/ год · Япония", 13, INK2, False)], align=PP_ALIGN.CENTER)
txt(s, Inches(7.3), Inches(5.15), Inches(5.0), Inches(1.2),
    [("Цифры берутся из админки («Страны и Университеты» + «Калькулятор стоимости обучения»)", 11, INK3, False),
     ("и обновляются без программиста.", 11, INK3, False)], align=PP_ALIGN.CENTER)
footer(s, "V · Калькулятор")

# ============================================================
# 12 · FORM
# ============================================================
s = add_slide()
kicker(s, "V · Сайт по экранам")
title(s, "Форма заявки — главный механизм лидов")
rows = [["Поле", "Обязательное"],
        ["Имя", ("да", BADL)], ["Телефон", ("да", BADL)], ["Откуда узнали о нас", ("да", BADL)],
        ["Email", ("нет", INK3)], ["Страна обучения", ("нет", INK3)],
        ["Желаемый вуз (с поиском)", ("нет", INK3)], ["Уровень, год, бюджет, язык", ("нет, раскрывается", INK3)]]
make_table(s, Inches(0.7), Inches(1.6), Inches(5.6), rows, [1.4, 1], size=11, row_h=Pt(24))
txt(s, Inches(0.7), Inches(4.95), Inches(5.6), Inches(0.7),
    [("Минимум обязательных полей = больше отправленных заявок.", 11, INK3, False),
     ("Остальное менеджер уточнит в разговоре.", 11, INK3, False)])
txt(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.4), [("Что после нажатия «Отправить»", 15, INK, True)])
steps_list(s, Inches(6.8), Inches(2.05), Inches(5.8),
           [("Экран «Спасибо»", "«Менеджер свяжется в течение 15 минут»"),
            ("Проверка на дубль", "Старый клиент? → комментарий к его карточке"),
            ("Назначение менеджера", "По правилам или очереди"),
            ("Push + Telegram", "Менеджер и РОП уведомлены"),
            ("Старт SLA-таймера", "3 часа на первый ответ")], size=12, gap=Inches(0.68))
card(s, Inches(0.7), Inches(5.8), Inches(11.9), Inches(1.05))
txt(s, Inches(0.95), Inches(5.95), Inches(11.4), Inches(0.8),
    [("ГДЕ ЕЩЁ РОЖДАЮТСЯ ЛИДЫ", 9, OKL, True),
     ("1) Менеджер создал вручную («+ Лид» — звонок или визит в офис)  ·  2) Партнёр привёл реферала  ·  3) Событийная страница — лид с меткой события", 11, INK2, False)])
footer(s, "V · Форма заявки")

# ============================================================
# 13 · ADMIN EDITS SITE
# ============================================================
s = add_slide()
kicker(s, "V · Сайт по экранам")
title(s, "Пример: как админ меняет сайт")
txt(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(0.4),
    [("Задача: поднять цену услуг с $500 до $700", 15, INK, True)])
steps_list(s, Inches(0.7), Inches(2.05), Inches(5.8),
           [("Открыть админку", "goglobal.kg/admin → пароль → главная"),
            ("«Калькулятор стоимости обучения»", "Кликнуть на заголовок секции"),
            ("Изменить число", "Поле «Стоимость услуг»: 500 → 700"),
            ("Нажать «Сохранить» сверху", "Появится «✓ сохранено»"),
            ("Проверить сайт", "goglobal.kg/#calculator — цифра обновилась")], size=12, gap=Inches(0.68))
txt(s, Inches(0.7), Inches(5.7), Inches(5.8), Inches(0.5),
    [("Итого: 5 минут, без программиста, без перезагрузок.", 12, OKL, False)])
rows = [["Задача", "Время"],
        ["Заменить главное фото сайта", "3 мин"], ["Добавить отзыв выпускника", "5 мин"],
        ["Добавить вуз в страну", "7 мин"], ["Создать страницу события", "10 мин"],
        ["Скрыть секцию сайта", "1 мин"], ["Обновить телефон в подвале", "2 мин"],
        ["Новый вопрос в FAQ", "3 мин"]]
make_table(s, Inches(6.8), Inches(1.6), Inches(5.8), rows, [2, 0.5], size=11, row_h=Pt(24))
txt(s, Inches(6.8), Inches(5.0), Inches(5.8), Inches(0.6),
    [("⚠ Не забывать кнопку «Сохранить» в шапке — без неё правки контента не применятся.", 11, WARNL, False)])
footer(s, "V · Правки сайта")

# ============================================================
# 14 · CRM WORKSPACE
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Рабочее место менеджера")
steps_list(s, Inches(0.7), Inches(1.7), Inches(5.8),
           [("Шапка", "Поиск · тумблер «в сети» · база знаний · «+ Лид»"),
            ("Сайдбар фильтров", "Статус, источник, страна, дата + пресеты"),
            ("«Мой день»", "Личный дашборд: встречи, горячие, задачи, суммы"),
            ("Список лидов", "5 видов отображения на выбор")], size=12.5, gap=Inches(0.72))
card(s, Inches(0.7), Inches(4.85), Inches(5.8), Inches(1.6), fill=RGBColor(0x2A, 0x1E, 0x0A), line_c=RGBColor(0x78, 0x50, 0x0F))
txt(s, Inches(0.95), Inches(5.0), Inches(5.3), Inches(1.3),
    [("КРИТИЧНАЯ ПРИВЫЧКА", 10, WARNL, True),
     ("Тумблер «🟢 В сети». Выключен — новые лиды НЕ приходят вам.", 11.5, INK, False),
     ("Утром включить · на обед выключить · вечером выключить.", 11.5, WARNL, True)])
# Мой день KPIs
txt(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.4), [("Виджет «Мой день»", 15, INK, True)])
mk = [("Встречи", "3", "назначено на сегодня"), ("Горячих", "5", "скоринг 60+ — приоритет"),
      ("Задач", "11 (2⚠)", "открытые (просрочены)"), ("В работе", "$67k", "сумма открытых сделок"),
      ("Прогноз", "$31k", "сумма × вероятность")]
for i, (l, v, sub) in enumerate(mk):
    x = Inches(6.8) + Inches(2.0) * (i % 3)
    y = Inches(2.0) + Inches(1.1) * (i // 3)
    kpi_box(s, x, y, Inches(1.85), Inches(0.95), l, v, sub, val_color=(OKL if l == "Прогноз" else INK))
txt(s, Inches(6.8), Inches(4.35), Inches(5.8), Inches(1.2),
    [("Виджет отвечает на вопрос «что делать прямо сейчас»:", 11.5, INK2, False),
     ("сначала просроченные задачи, потом горячие лиды, потом встречи.", 11.5, INK2, False)])
footer(s, "VI · CRM")

# ============================================================
# 15 · FIVE VIEWS
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Пять видов списка")
rows = [["Вид", "Что показывает", "Когда удобен", "Фишка"],
        [("🪟 Карточки", INK, True), "Сетка крупных карточек", "Общий обзор глазами", "Метки и суммы видны сразу"],
        [("📋 Таблица", INK, True), "Компактные строки, 9 колонок", "Много лидов, нужна плотность", "Скан + массовый выбор"],
        [("🎯 Воронка статусов", INK, True), "Канбан-колонки по статусам", "Управление потоком", ("Перетаскивание = смена статуса", OKL, True)],
        [("🎓 Этапы клиентов", INK, True), "Канбан только для Won", "Ведение после контракта", "Виза, оплаты, отъезд"],
        [("📅 Календарь", INK, True), "Встречи: список/неделя/месяц", "Планирование недели", "Клик по встрече = карточка"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [1.0, 1.2, 1.1, 1.3], size=10.5, row_h=Pt(27))
card(s, Inches(0.7), Inches(4.6), Inches(5.85), Inches(2.05))
txt(s, Inches(0.95), Inches(4.75), Inches(5.4), Inches(1.8),
    [("ВЫБОР ЗАПОМИНАЕТСЯ", 10, ACCENT, True),
     ("Каждый менеджер выбирает свой вид. Завтра CRM откроется в том же виде, с теми же фильтрами.", 11.5, INK2, False),
     ("Настройки индивидуальны.", 11.5, INK2, False)])
card(s, Inches(6.75), Inches(4.6), Inches(5.85), Inches(2.05))
txt(s, Inches(7.0), Inches(4.75), Inches(5.4), Inches(1.8),
    [("ФИЛЬТРЫ-ПРЕСЕТЫ", 10, ACCENT, True),
     ("Частые наборы («Горячие японцы», «Сентябрь-2026») сохраняются одной кнопкой", 11.5, INK2, False),
     ("и применяются в один клик.", 11.5, INK2, False)])
footer(s, "VI · Виды списка")

# ============================================================
# 16 · LEAD CARD
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Карточка клиента — 7 вкладок")
rows = [["Вкладка", "Что внутри", "Главное действие"],
        [("Обзор", INK, True), "Выпадашка статуса с подсказками, степпер этапов, контакты, метки, передача", "Сменить статус лида"],
        [("💰 Сделка", INK, True), "Сумма, валюта, вероятность 0–100%", "Зафиксировать бюджет"],
        [("📋 Задачи", INK, True), "Напоминания с дедлайнами", "«Перезвонить в среду 14:00»"],
        [("📎 Файлы", INK, True), "Паспорт, диплом, IELTS, контракт (до 25 МБ)", "Загрузить документ"],
        [("💬 Чат", INK, True), "Внутренние комментарии команды", "Заметка для коллег"],
        [("🕒 Аудит", INK, True), "Полная история: кто и что менял", "Разбор спорной ситуации"],
        [("🔗 Связанные", INK, True), "Лиды с тем же телефоном/email", "Прошлые обращения"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [0.8, 1.7, 1.2], size=10.5, row_h=Pt(24))
card(s, Inches(0.7), Inches(4.95), Inches(5.85), Inches(1.85))
txt(s, Inches(0.95), Inches(5.1), Inches(5.4), Inches(1.6),
    [("БЫСТРЫЕ ДЕЙСТВИЯ", 10, ACCENT, True),
     ("💬 WhatsApp · 📞 Позвонить · ✉ Email · 📨 Шаблон", 12, INK, True),
     ("«Шаблон» — готовые сообщения с автоподстановкой имени; WhatsApp откроется с готовым текстом.", 11, INK2, False)])
card(s, Inches(6.75), Inches(4.95), Inches(5.85), Inches(1.85))
txt(s, Inches(7.0), Inches(5.1), Inches(5.4), Inches(1.6),
    [("КАЖДОЕ КАСАНИЕ — НА СЧЕТУ", 10, ACCENT, True),
     ("Клик по WhatsApp/телефону/email = «касание клиента» в статистике.", 11, INK2, False),
     ("Скоринг 0–100 оценивает «горячесть» лида: 60+ = горячий, в первую очередь.", 11, INK2, False)])
footer(s, "VI · Карточка")

# ============================================================
# 17 · STATUSES
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Статусы: две воронки")
rows1 = [["Статус", "Что значит"],
         [("Новый", SKYL, True), "Только поступил, ещё не тронут"],
         [("В работе", WARNL, True), "Первый контакт состоялся"],
         [("Перезвонить", WARNL, True), "Клиент попросил позже"],
         [("Не ответил", WARNL, True), "Не берёт трубку"],
         [("Подойдёт в офис", WARNL, True), "Встреча (система спросит дату)"],
         [("Закрыт ✅ (won)", OKL, True), "Контракт подписан"],
         [("Отказ ❌ (lost)", BADL, True), "Система спросит причину"]]
txt(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(0.6),
    [("Воронка 1 · Обработка лида", 15, INK, True), ("От первого контакта до решения", 11, INK3, False)])
make_table(s, Inches(0.7), Inches(2.15), Inches(5.8), rows1, [0.8, 1.4], size=10.5, row_h=Pt(23))
rows2 = [["Этап", "Суть"],
         ["1. 📜 Контракт подписан", "Договор"], ["2. 💰 Оплата 1", "Предоплата"],
         ["3. 📂 Сбор документов", "Сбор и перевод"], ["4. 🇬🇧 Языковой экзамен", "IELTS / TOPIK / JLPT"],
         ["5. 🎤 Собеседование", "Интервью с вузом"], ["6. 🎓 Зачисление", "Вуз принял"],
         ["7. 🛂 Виза", "Оформление визы"], ["8. 💵 Окончательная оплата", "Финальный платёж"],
         ["9. ✈️ Отъезд / прибытие", "Клиент уехал учиться"]]
txt(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.6),
    [("Воронка 2 · Этапы клиента", 15, INK, True), ("После Won — до отъезда на учёбу", 11, INK3, False)])
make_table(s, Inches(6.8), Inches(2.15), Inches(5.8), rows2, [0.9, 1.3], size=10.5, row_h=Pt(21))
card(s, Inches(0.7), Inches(6.05), Inches(5.8), Inches(0.95))
txt(s, Inches(0.95), Inches(6.18), Inches(5.4), Inches(0.75),
    [("КАК СВЯЗАНЫ ВОРОНКИ", 9, ACCENT, True),
     ("Этапы заблокированы до победы. «Закрыт ✅» автоматически открывает «Контракт подписан»; откат статуса снимает этап.", 10.5, INK2, False)])
footer(s, "VI · Статусы")

# ============================================================
# 18 · EFFICIENCY TOOLS
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Инструменты эффективности")
tools = [("📥 Inbox 0", "Видны только лиды, требующие действий. Закрыл всё — список пуст."),
         ("💾 Свои фильтры", "Настроил раз — применяешь одним кликом каждый день."),
         ("☑️ Массовые действия", "Выделил 10 лидов → сменил статус всем сразу."),
         ("📨 Шаблоны ответов", "Готовые тексты, имя подставится само."),
         ("🔄 Передача лида", "Отпуск? Передал коллеге. 10 минут на «принять»."),
         ("📲 Push на телефон", "CRM как приложение: новый лид — уведомление сразу.")]
for i, (name, desc) in enumerate(tools):
    x = Inches(0.7) + Inches(4.05) * (i % 3)
    y = Inches(1.55) + Inches(1.25) * (i // 3)
    card(s, x, y, Inches(3.85), Inches(1.1))
    txt(s, x + Inches(0.2), y + Inches(0.1), Inches(3.5), Inches(0.95),
        [(name, 13, INK, True), (desc, 10, INK2, False)])
rows = [["Инструмент", "Без него", "С ним", "Экономия"],
        ["Шаблоны ответов", "Печатать каждое ~2 мин", "2 клика ~5 сек", ("~30 мин", OKL, True)],
        ["Inbox 0", "Листать все 50 лидов", "Видеть 8 актуальных", ("~20 мин", OKL, True)],
        ["Массовые действия", "Открывать каждый лид", "Выделить — применить", ("~15 мин", OKL, True)],
        ["Сохранённые фильтры", "Настраивать заново", "1 клик", ("~10 мин", OKL, True)],
        [("Итого в день", INK, True), "", "", ("~1 ч 15 мин", OKL, True)]]
make_table(s, Inches(0.7), Inches(4.25), Inches(11.9), rows, [1.0, 1.2, 1.1, 0.6], size=10.5, row_h=Pt(23))
footer(s, "VI · Инструменты")

# ============================================================
# 19 · SAFETY NETS
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Шесть защитных сеток — работают сами")
rows = [["Защита", "Когда срабатывает", "Что делает"],
        [("⏰ SLA-таймер", INK, True), "Нет ответа 3 часа", "Лид краснеет, РОПу сигнал, минус в статистику"],
        [("🔁 Дедупликация", INK, True), "Заявка с тем же телефоном", "Комментарий к старому лиду, не дубль"],
        [("🤖 Авто-сценарии", INK, True), "Лид завис в статусе N часов", "Задача менеджеру / сигнал (настраивается)"],
        [("🌅 Утренний отчёт", INK, True), "Каждый день в 09:00", "Сводка РОПу в Telegram"],
        [("📊 Аудит", INK, True), "Каждое действие", "Кто/что/когда — споры решаются за минуту"],
        [("💾 Авто-бэкап", INK, True), "Ежедневно", "Копия всех данных"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [1.0, 1.1, 1.7], size=10.5, row_h=Pt(25))
card(s, Inches(0.7), Inches(4.55), Inches(5.85), Inches(2.1), fill=RGBColor(0x0A, 0x13, 0x22))
txt(s, Inches(0.95), Inches(4.7), Inches(5.4), Inches(1.85),
    [("ПРИМЕР АВТО-СЦЕНАРИЯ", 10, ACCENT, True),
     ("ЕСЛИ лид в статусе «Перезвонить»", 11.5, INK2, False),
     ("       дольше 48 часов", 11.5, INK2, False),
     ("ТО создать задачу менеджеру:", 11.5, INK2, False),
     ("       «Срочно перезвонить — клиент ждёт»", 11.5, WARNL, False),
     ("Настраивается в админке без программиста.", 10, INK3, False)])
card(s, Inches(6.75), Inches(4.55), Inches(5.85), Inches(2.1))
txt(s, Inches(7.0), Inches(4.7), Inches(5.4), Inches(0.35), [("НА ПРАКТИКЕ", 10, OKL, True)])
bullet_list(s, Inches(7.0), Inches(5.05), Inches(5.4), Inches(1.5),
            ["Ни один лид не «теряется в столе»", "Отпуск — система подскажет передать дела",
             "Повторный клиент — вся история на месте", "Спор «кто упустил» решается фактами"], size=10.5)
footer(s, "VI · Защитные сетки")

# ============================================================
# 20 · CRM EXTRAS + LOSS REASONS
# ============================================================
s = add_slide()
kicker(s, "VI · CRM продавца")
title(s, "Ещё четыре возможности")
extras = [("📖 База знаний", "РОП пишет статьи («Виза в Японию», «Возражение дорого») — менеджер открывает из CRM. Новички учатся сами."),
          ("🎟 Метки клиентов", "Цветные ярлыки: Hot, VIP, Грант, Повторный. Видны сразу, фильтруются."),
          ("📅 Календарь встреч", "Статус «Подойдёт в офис» + дата → встреча сама в календаре."),
          ("🚪 Причины отказов", "При Lost система просит причину: дорого / передумал / конкурент. РОП видит сводку.")]
for i, (name, desc) in enumerate(extras):
    x = Inches(0.7) + Inches(6.1) * (i % 2)
    y = Inches(1.55) + Inches(1.3) * (i // 2)
    card(s, x, y, Inches(5.9), Inches(1.15))
    txt(s, x + Inches(0.22), y + Inches(0.12), Inches(5.5), Inches(0.95),
        [(name, 13.5, INK, True), (desc, 10.5, INK2, False)])
# loss reasons bars
txt(s, Inches(0.7), Inches(4.35), Inches(11), Inches(0.4),
    [("Пример сводки причин отказов (за квартал)", 14, INK, True)])
reasons = [("Дорого", 42, "$84k", BAD), ("Не отвечает", 25, "$37k", WARN),
           ("Передумал", 18, "$29k", RGBColor(0x63, 0x66, 0xF1)), ("Конкурент", 10, "$18k", ACCENT),
           ("Другое", 5, "$8k", INK3)]
for i, (nm, pct, amt, color) in enumerate(reasons):
    y = Inches(4.85) + Inches(0.38) * i
    txt(s, Inches(0.7), y, Inches(1.6), Inches(0.32), [(nm, 11, INK, True)])
    track = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.4), y + Inches(0.04), Inches(8.0), Inches(0.22))
    track.adjustments[0] = 0.5
    track.fill.solid(); track.fill.fore_color.rgb = BG2; track.line.fill.background(); track.shadow.inherit = False
    seg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.4), y + Inches(0.04), Inches(8.0 * pct / 100), Inches(0.22))
    seg.adjustments[0] = 0.5
    seg.fill.solid(); seg.fill.fore_color.rgb = color; seg.line.fill.background(); seg.shadow.inherit = False
    txt(s, Inches(10.6), y, Inches(2.0), Inches(0.32), [(f"{pct}% · {amt}", 10.5, INK2, False)], align=PP_ALIGN.RIGHT)
txt(s, Inches(0.7), Inches(6.8), Inches(11.9), Inches(0.4),
    [("Вывод РОПа: 42% отказов «дорого» → добавить рассрочку и гранты в скрипты продаж.", 11, INK3, False)])
footer(s, "VI · Дополнительно")

# ============================================================
# 21 · ADMIN STRUCTURE
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "30+ разделов в трёх группах")
groups = [("📊 Аналитика · 11",
           ["Здоровье продаж (светофор)", "Посещения сайта", "Дашборд CRM", "Leaderboard",
            "1-на-1 отчёт", "Время в статусах", "Скорость отклика", "Heatmap входящих",
            "Причины отказов", "Когортный анализ", "ROI по источникам"]),
          ("⚙️ Операции · 12",
           ["Менеджеры", "Авто-распределение", "Статусы воронки", "Метки клиентов",
            "Шаблоны ответов", "Авто-сценарии", "Причины отказов", "База знаний",
            "Комиссии и выплаты", "Журнал аудита", "Утилиты и бэкапы", "Все лиды"]),
          ("🌍 Контент сайта · 10",
           ["Видимость блоков", "Главный экран (Hero)", "О компании", "Страны и Университеты",
            "Партнёрские университеты", "Калькулятор", "Отзывы", "FAQ",
            "Контакты и график", "События"])]
for i, (name, items) in enumerate(groups):
    x = Inches(0.7) + Inches(4.05) * i
    card(s, x, Inches(1.55), Inches(3.85), Inches(4.3))
    txt(s, x + Inches(0.22), Inches(1.7), Inches(3.5), Inches(0.4), [(name, 14.5, INK, True)])
    bullet_list(s, x + Inches(0.22), Inches(2.2), Inches(3.5), Inches(3.6), items, size=10, gap_para=False)
card(s, Inches(0.7), Inches(6.05), Inches(11.9), Inches(0.95))
txt(s, Inches(0.95), Inches(6.18), Inches(11.4), Inches(0.75),
    [("БЕРЕЖНЫЙ РЕЖИМ", 9, OKL, True),
     ("Все разделы свёрнуты; раздел грузит данные только когда его открыли — админка лёгкая даже на слабых ноутбуках.", 11, INK2, False)])
footer(s, "VII · Структура админки")

# ============================================================
# 22 · TRAFFIC LIGHT
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "«Светофор» — состояние продаж за 2 секунды")
rows = [["Цвет", "Условия", "Действие РОПа"],
        [("🟢 Норма", OKL, True), "SLA ≥ 80% · конверсия ≥ 15% · просрочек < 6", "Работаем дальше"],
        [("🟡 Внимание", WARNL, True), "SLA 60–80% или конверсия 10–15%", "Открыть детали, найти узкое место"],
        [("🔴 Срочно", BADL, True), "SLA < 60% или конверсия < 10% или просрочек > 5", "Разбираться немедленно"]]
make_table(s, Inches(0.7), Inches(1.6), Inches(11.9), rows, [0.7, 1.9, 1.1], size=11.5, row_h=Pt(30))
# mock traffic light
card(s, Inches(0.7), Inches(3.55), Inches(11.9), Inches(1.45), fill=RGBColor(0x0A, 0x1F, 0x18), line_c=RGBColor(0x10, 0xB9, 0x81))
txt(s, Inches(1.0), Inches(3.78), Inches(7.0), Inches(1.0),
    [("🟢  ЗДОРОВЬЕ ПРОДАЖ", 13, OKL, True), ("ВСЁ В ПОРЯДКЕ", 19, WHITE, True)])
mk = [("SLA", "87%", OKL), ("Конверсия", "22%", OKL), ("Просрочки", "2", INK), ("Зависшие", "0", INK)]
for i, (l, v, c) in enumerate(mk):
    txt(s, Inches(7.3) + Inches(1.3) * i, Inches(3.75), Inches(1.25), Inches(1.0),
        [(l.upper(), 9, INK3, True), (v, 20, c, True)])
card(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.3))
txt(s, Inches(0.95), Inches(5.45), Inches(11.4), Inches(1.05),
    [("ЕСЛИ ЖЁЛТЫЙ ИЛИ КРАСНЫЙ", 10, WARNL, True),
     ("Под светофором — список причин: «SLA упал из-за Дамира — офлайн 24ч». Не нужно копать отчёты —", 11.5, INK2, False),
     ("система сама говорит, где проблема и у кого.", 11.5, INK2, False)])
footer(s, "VII · Светофор")

# ============================================================
# 23 · LEADERBOARD
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "Leaderboard — рейтинг команды")
rows = [["#", "Менеджер", "Лидов", "Won", "Конверсия", "Выручка", "Отклик", "Касаний/д"],
        ["🥇", ("Гулжан Тен", INK, True), "64", "15", ("23%", OKL), ("$32 100", OKL), ("18 мин", OKL), "31"],
        ["🥈", ("Анна Иванова", INK, True), "71", "18", ("25%", OKL), ("$28 500", OKL), ("12 мин", OKL), "28"],
        ["🥉", ("Бек Бекович", INK, True), "55", "11", "20%", "$19 200", ("32 мин", WARNL), "22"],
        ["4", ("Дамир С. ⚠", BADL, True), "38", "6", ("16%", BADL), ("$8 400", BADL), ("94 мин", BADL), ("9", BADL)]]
make_table(s, Inches(0.7), Inches(1.6), Inches(11.9), rows, [0.3, 1.3, 0.55, 0.45, 0.8, 0.85, 0.7, 0.8], size=11, row_h=Pt(28))
card(s, Inches(0.7), Inches(4.05), Inches(5.85), Inches(2.45))
txt(s, Inches(0.95), Inches(4.2), Inches(5.4), Inches(0.35), [("КАК ЧИТАТЬ", 10, ACCENT, True)])
bullet_list(s, Inches(0.95), Inches(4.55), Inches(5.4), Inches(1.85),
            [[("Конверсия", INK, True), (" — главный показатель качества", INK2, False)],
             [("Отклик", INK, True), (" — скорость 1-го ответа: норма < 30 мин", INK2, False)],
             [("Касания/день", INK, True), (" — реально ли работает с базой", INK2, False)],
             [("Всё низкое сразу", INK, True), (" = повод для 1-на-1", INK2, False)]], size=11)
card(s, Inches(6.75), Inches(4.05), Inches(5.85), Inches(2.45))
txt(s, Inches(7.0), Inches(4.2), Inches(5.4), Inches(2.2),
    [("СВЯЗКА С 1-НА-1 ОТЧЁТОМ", 10, ACCENT, True),
     ("Увидел проблему → «1-на-1 отчёт» → выбрал менеджера → система собрала недельную сводку:", 11.5, INK2, False),
     ("KPI с динамикой, топ-лиды, зависшие, касания по каналам.", 11.5, INK2, False),
     ("Подготовка к встрече: 0 минут.", 12, OKL, True)])
footer(s, "VII · Leaderboard")

# ============================================================
# 24 · ANALYTICS WIDGETS
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "Остальные виджеты аналитики")
rows = [["Виджет", "Вопрос", "Пример вывода"],
        [("⏱ Время в статусах", INK, True), "Где зависают лиды?", "«В Перезвонить сидят 5 дней — забывают»"],
        [("⏲ Скорость отклика", INK, True), "Кто быстрый, кто тянет?", "«Анна — 12 мин, Дамир — 94 мин»"],
        [("🌡 Heatmap входящих", INK, True), "Когда приходят заявки?", "«Пик — вс 20:00, все офлайн → дежурства»"],
        [("📆 Когортный анализ", INK, True), "Конверсия растёт или падает?", "«Май конвертится хуже марта на 7%»"],
        [("💸 ROI источников", INK, True), "Какая реклама окупается?", "«Instagram: $500 → $12k. TikTok: $800 → $2k»"],
        [("📈 Дашборд CRM", INK, True), "Общая картина за период?", "Лиды по дням, источникам + прогноз выручки"],
        [("👁 Посещения сайта", INK, True), "Сколько людей заходит?", "Динамика трафика по дням"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [1.0, 1.0, 1.6], size=10.5, row_h=Pt(25))
card(s, Inches(0.7), Inches(4.95), Inches(5.85), Inches(1.9))
txt(s, Inches(0.95), Inches(5.1), Inches(5.4), Inches(1.65),
    [("ПРИМЕР ВЫВОДА ИЗ HEATMAP", 10, ACCENT, True),
     ("Воскресенье 18:00–22:00 — пик заявок (17/час), но вся команда офлайн.", 11.5, INK2, False),
     ("Решение: дежурный менеджер по воскресеньям → +15–20 спасённых лидов в месяц.", 11.5, INK2, False)])
card(s, Inches(6.75), Inches(4.95), Inches(5.85), Inches(1.9))
txt(s, Inches(7.0), Inches(5.1), Inches(5.4), Inches(1.65),
    [("ЭКСПОРТ В EXCEL", 10, ACCENT, True),
     ("Любой период — одной кнопкой из Дашборда: все лиды с полями, статусами, источниками, суммами.", 11.5, INK2, False),
     ("Для бухгалтерии и инвесторов.", 11.5, INK2, False)])
footer(s, "VII · Аналитика")

# ============================================================
# 25 · TEAM MGMT
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "Управление командой")
txt(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(0.4), [("Создать менеджера — 2 минуты", 15, INK, True)])
steps_list(s, Inches(0.7), Inches(2.0), Inches(5.8),
           [("Админка → «Менеджеры»", "Раскрыть раздел"),
            ("Заполнить форму", "Имя, логин (латиницей), Telegram-тег"),
            ("Выбрать роль", "Менеджер / РОП / Партнёр"),
            ("«+ Создать»", "Передать сотруднику логин и пароль")], size=12, gap=Inches(0.62))
card(s, Inches(6.8), Inches(1.55), Inches(5.8), Inches(2.9), fill=RGBColor(0x22, 0x10, 0x12), line_c=RGBColor(0x7F, 0x1D, 0x1D))
txt(s, Inches(7.05), Inches(1.7), Inches(5.3), Inches(2.6),
    [("УВОЛИТЬ = АРХИВИРОВАТЬ, НЕ УДАЛЯТЬ", 11, BADL, True),
     ("Кнопка «Архивировать»:", 11.5, INK2, False),
     ("• система спросит, кому передать открытые лиды", 11.5, INK2, False),
     ("• доступ закрывается сразу", 11.5, INK2, False),
     ("• история работы сохраняется", 11.5, INK2, False),
     ("Никогда не удаляйте насовсем — сломается статистика прошлых периодов.", 11.5, BADL, True)])
rows = [["Правило (пример)", "Результат"],
        ["Страна = Япония → Анна", "«Японские» лиды идут Анне (спец по Японии)"],
        ["Источник = Instagram → Гулжан", "Лиды из Instagram — Гулжан (ведёт соцсети)"],
        ["Правил нет / не подошло", "Очередь по кругу между «🟢 в сети»"],
        ["Менеджер из правила офлайн", "Лид в общую очередь — не теряется"]]
txt(s, Inches(0.7), Inches(4.75), Inches(11), Inches(0.4), [("Авто-распределение: кому падают лиды", 14, INK, True)])
make_table(s, Inches(0.7), Inches(5.2), Inches(11.9), rows, [1, 1.5], size=11, row_h=Pt(24))
footer(s, "VII · Команда")

# ============================================================
# 26 · COMMISSIONS + AUDIT
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "Комиссии и журнал аудита")
txt(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(0.4), [("💵 Комиссии и выплаты", 15, INK, True)])
rows = [["Правило", "Пример"],
        ["Процент от сделки", "5% от суммы Won-контракта"],
        ["Фикс за сделку", "$100 за каждый Won"],
        ["Для партнёров", "3% от приведённого реферала"]]
make_table(s, Inches(0.7), Inches(2.0), Inches(5.8), rows, [1, 1.3], size=11, row_h=Pt(26))
txt(s, Inches(0.7), Inches(3.95), Inches(5.8), Inches(0.9),
    [("Система сама считает сумму к выплате по каждому менеджеру за период.", 11.5, INK2, False),
     ("РОП открывает раздел — видит готовые цифры.", 11.5, INK2, False)])
txt(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.4), [("🕒 Журнал аудита", 15, INK, True)])
card(s, Inches(6.8), Inches(2.0), Inches(5.8), Inches(2.6), fill=RGBColor(0x0A, 0x13, 0x22))
txt(s, Inches(7.05), Inches(2.2), Inches(5.3), Inches(2.3),
    [("14:02  Анна сменила статус #1247", 11, INK2, False),
     ("           «Новый → В работе»", 11, INK3, False),
     ("14:05  Анна загрузила «passport.pdf»", 11, INK2, False),
     ("14:31  Бек (РОП) переназначил #1198", 11, INK2, False),
     ("           «Дамир → Гулжан»", 11, INK3, False),
     ("15:10  Система: SLA-просрочка #1201", 11, BADL, False)])
card(s, Inches(0.7), Inches(5.1), Inches(11.9), Inches(1.3))
txt(s, Inches(0.95), Inches(5.27), Inches(11.4), Inches(1.0),
    [("ЗАЧЕМ", 10, ACCENT, True),
     ("Спор «кто упустил клиента» решается за минуту — открыл аудит лида и видишь всю историю действий каждого сотрудника.", 12, INK2, False)])
footer(s, "VII · Комиссии и аудит")

# ============================================================
# 27 · DATA SAFETY
# ============================================================
s = add_slide()
kicker(s, "VII · Админка и аналитика")
title(s, "Сохранность данных")
saf = [("💾 Авто-бэкап", "Ежедневная копия базы автоматически, на стороне хостинга. Сгорел сервер — данные восстановимы."),
       ("📦 Ручной экспорт", "«Утилиты и бэкапы»: одна кнопка — полный архив. Раз в месяц — в облако компании."),
       ("📊 Excel-выгрузки", "Все лиды за период — в Excel из Дашборда. Для бухгалтерии и отчётов.")]
for i, (name, desc) in enumerate(saf):
    x = Inches(0.7) + Inches(4.05) * i
    card(s, x, Inches(1.55), Inches(3.85), Inches(1.6))
    txt(s, x + Inches(0.2), Inches(1.7), Inches(3.45), Inches(1.35),
        [(name, 14, INK, True), (desc, 10.5, INK2, False)])
rows = [["Что", "Как устроено"],
        ["Пароль админки", "Один общий, у админа. Менять раз в квартал"],
        ["Пароли менеджеров", "Личный логин у каждого; пароль выдаёт админ"],
        ["Вход в CRM", "Сессия 12 часов — хватает на весь рабочий день"],
        ["Уволенный сотрудник", "Архивация мгновенно закрывает доступ"],
        ["Документы клиентов", "На сервере платформы, доступ только из CRM"]]
txt(s, Inches(0.7), Inches(3.45), Inches(11), Inches(0.4), [("Безопасность доступов", 14, INK, True)])
make_table(s, Inches(0.7), Inches(3.9), Inches(11.9), rows, [0.8, 1.8], size=11, row_h=Pt(25))
footer(s, "VII · Данные")

# ============================================================
# 28 · CLIENT JOURNEY TIMELINE
# ============================================================
s = add_slide()
kicker(s, "VIII · Путь клиента")
title(s, "От заявки до отъезда — таймлайн")
events = [("День 0 · 📨 Заявка на сайте", "Айгуль заполнила форму → лид #1247 «Новый» упал Анне (правило «Япония → Анна») → push + SLA 3 ч"),
          ("День 0, +20 мин · 💬 Первый контакт", "Кнопка WhatsApp → шаблон «Приветствие» → разговор → статус «В работе»"),
          ("День 2 · 🤝 Встреча в офисе", "Статус «Подойдёт в офис» + дата → встреча в календаре → бюджет $15k"),
          ("День 3 · 📎 Документы и сделка", "Паспорт и аттестат — в «Файлы». Сделка: $15 000, вероятность 70%, метка Hot"),
          ("День 7 · 🎉 Контракт — Won", "Лид перешёл во вторую воронку «Этапы клиента» → +1 в Leaderboard Анны"),
          ("Месяцы 1–6 · 🎓 Сопровождение", "Оплата 1 → Документы → Языковой экзамен → Собеседование → Зачисление в Toyo University → Виза → Оплата → Отъезд"),
          ("Месяц 6 · ✈️ Отъезд", "Клиентка в Японии. Карточка остаётся в истории — для отзывов и рефералов")]
# vertical line
vline = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(1.7), Pt(2), Inches(5.1))
vline.fill.solid(); vline.fill.fore_color.rgb = LINE; vline.line.fill.background(); vline.shadow.inherit = False
for i, (head, body) in enumerate(events):
    y = Inches(1.65) + Inches(0.74) * i
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.83), y + Inches(0.05), Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = BG
    dot.line.color.rgb = ACCENT; dot.line.width = Pt(1.75); dot.shadow.inherit = False
    txt(s, Inches(1.35), y, Inches(11.2), Inches(0.7),
        [(head, 12.5, INK, True), (body, 11, INK2, False)])
footer(s, "VIII · Путь клиента")

# ============================================================
# 29 · JOURNEY METRICS
# ============================================================
s = add_slide()
kicker(s, "VIII · Путь клиента")
title(s, "Что фиксируется на каждом шаге")
rows = [["Шаг", "Что записала система", "Где видно потом"],
        ["Заявка пришла", "Время, источник, UTM", "Heatmap · ROI источников"],
        ["Первый ответ", "Время реакции менеджера", "Скорость отклика · SLA"],
        ["Каждый звонок/WhatsApp", "Касание: канал + время", "Leaderboard (касания/день)"],
        ["Смена статуса", "Кто, когда, из какого в какой", "Время в статусах · Аудит"],
        ["Сумма сделки", "$ и вероятность", "Прогноз выручки в Дашборде"],
        ["Won/Lost", "Результат + причина отказа", "Конверсия · Причины отказов"],
        ["Этапы после Won", "Прогресс сопровождения", "Канбан «Этапы клиентов»"]]
make_table(s, Inches(0.7), Inches(1.55), Inches(11.9), rows, [1.0, 1.2, 1.2], size=11, row_h=Pt(25))
card(s, Inches(0.7), Inches(5.0), Inches(5.85), Inches(1.8))
txt(s, Inches(0.95), Inches(5.15), Inches(5.4), Inches(1.55),
    [("ГЛАВНАЯ МЫСЛЬ", 10, ACCENT, True),
     ("Менеджер просто работает — звонит, пишет, двигает статусы. Вся аналитика собирается сама.", 11.5, INK2, False),
     ("Никто ничего не заполняет «для отчёта».", 11.5, INK, True)])
card(s, Inches(6.75), Inches(5.0), Inches(5.85), Inches(1.8))
txt(s, Inches(7.0), Inches(5.15), Inches(5.4), Inches(1.55),
    [("ДЛЯ КЛИЕНТА ЭТО ЗНАЧИТ", 10, OKL, True),
     ("Его не «теряют»: любой менеджер откроет карточку и за 30 секунд поймёт всю историю —", 11.5, INK2, False),
     ("кто что обещал, какие документы есть, что дальше.", 11.5, INK2, False)])
footer(s, "VIII · Метрики пути")

# ============================================================
# 30 · STORY 1 (Anna)
# ============================================================
s = add_slide()
kicker(s, "IX · Истории из практики")
title(s, "История 1 · Утро менеджера Анны")
card(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(1.0))
txt(s, Inches(0.95), Inches(1.62), Inches(5.4), Inches(0.8),
    [[("Анна Иванова", 13, INK, True), ("  · менеджер · 18 won за неделю", 11, INK3, False)],
     ("«Я не думаю утром, что делать первым. CRM уже подумала за меня».", 11.5, INK2, False)])
steps_list(s, Inches(0.7), Inches(2.8), Inches(5.8),
           [("09:00 · Открыла /lidy", "«Мой день»: 3 встречи, 5 горячих, 2 просрочки"),
            ("09:02 · Inbox 0", "Из 47 лидов осталось 8 требующих действия"),
            ("09:05 · Закрыла просрочки", "2 звонка — задачи закрыты"),
            ("09:20 · Новые лиды", "3 шаблона «Приветствие» в WhatsApp"),
            ("10:00 · Уехала на встречу", "Тумблер «не в сети» — лиды идут коллегам")], size=11.5, gap=Inches(0.64))
ik = [("Обработано", "12", INK), ("Касаний", "28", INK), ("Won", "2 · $24.5k", OKL), ("Inbox 0", "✓ пуст", OKL)]
for i, (l, v, c) in enumerate(ik):
    x = Inches(6.8) + Inches(2.95) * (i % 2)
    y = Inches(1.6) + Inches(1.1) * (i // 2)
    kpi_box(s, x, y, Inches(2.8), Inches(0.95), l, v, val_color=c)
card(s, Inches(6.8), Inches(3.95), Inches(5.8), Inches(2.3))
txt(s, Inches(7.05), Inches(4.1), Inches(5.3), Inches(2.05),
    [("ПОЧЕМУ ЭТО РАБОТАЕТ", 10, OKL, True),
     ("Система всегда показывает следующий шаг.", 12, INK, True),
     ("Меньше решений «что делать» → меньше усталость → выше конверсия.", 11.5, INK2, False),
     ("Все 28 касаний зачтены в статистику автоматически.", 11.5, INK2, False)])
footer(s, "IX · История 1")

# ============================================================
# 31 · STORY 2 (ROP)
# ============================================================
s = add_slide()
kicker(s, "IX · Истории из практики")
title(s, "История 2 · Понедельник РОПа Бека")
card(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(1.5), fill=RGBColor(0x00, 0x1B, 0x3D), line_c=RGBColor(0x0C, 0x4A, 0x6E))
txt(s, Inches(0.95), Inches(1.62), Inches(5.4), Inches(1.3),
    [("📊 GoGlobal · утренний отчёт", 12, WHITE, True),
     ("За выходные: 23 лида, 3 Won ($12 500), 1 Lost", 11.5, INK2, False),
     ("⚠️ Просрочено SLA: 2 · 📋 Просроченных задач: 4", 11.5, INK2, False),
     ("⚪ Офлайн >24ч: Дамир", 11.5, WARNL, False)])
steps_list(s, Inches(0.7), Inches(3.3), Inches(5.8),
           [("09:05 · Светофор: 🟡", "Причина — SLA 72% из-за Дамира"),
            ("09:10 · Leaderboard", "У Дамира 94 мин отклик — кандидат на 1-на-1"),
            ("09:20 · Звонок Дамиру", "Он в больнице — нужно передать лиды"),
            ("09:25 · Массовая передача", "Фильтр «Дамир + Hot» → 5 лидов → Анне за 30 сек")], size=11.5, gap=Inches(0.66))
card(s, Inches(6.8), Inches(1.55), Inches(5.8), Inches(2.2), fill=RGBColor(0x0A, 0x13, 0x22))
txt(s, Inches(7.05), Inches(1.7), Inches(5.3), Inches(1.95),
    [("АВТО-СЦЕНАРИЙ НА ВРЕМЯ БОЛЕЗНИ", 10, ACCENT, True),
     ("ЕСЛИ лид Дамира в статусе «Перезвонить» > 24ч", 11.5, INK2, False),
     ("ТО уведомить РОПа в Telegram", 11.5, INK2, False),
     ("Ничего не «протухнет» незаметно, пока Дамир болеет.", 11.5, OKL, False)])
card(s, Inches(6.8), Inches(4.0), Inches(5.8), Inches(2.25))
txt(s, Inches(7.05), Inches(4.15), Inches(5.3), Inches(0.35), [("ИТОГ ПОНЕДЕЛЬНИКА", 10, OKL, True)])
bullet_list(s, Inches(7.05), Inches(4.5), Inches(5.3), Inches(1.7),
            ["Проблема найдена за 5 минут (не в пятницу)", "5 горячих лидов спасены передачей",
             "Контроль на время болезни — автоматический", "Ни одной таблицы Excel не открыто"], size=11)
footer(s, "IX · История 2")

# ============================================================
# 32 · STORY 3 (Partner)
# ============================================================
s = add_slide()
kicker(s, "IX · Истории из практики")
title(s, "История 3 · Партнёр приводит реферала")
card(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(1.1))
txt(s, Inches(0.95), Inches(1.62), Inches(5.4), Inches(0.9),
    [[("Нурлан", 13, INK, True), ("  · партнёр · преподаватель английского", 11, INK3, False)],
     ("«Просто завожу учеников в систему — и получаю процент».", 11.5, INK2, False)])
steps_list(s, Inches(0.7), Inches(2.85), Inches(5.8),
           [("Ученица спросила про Корею", "«Передам тебя проверенному агентству»"),
            ("Вошёл в /lidy под своим логином", "Роль «партнёр» — видит только своих"),
            ("«+ Лид»: имя и телефон", "Лид помечен «от партнёра Нурлана»"),
            ("Сделка Won $12 000", "Комиссия посчиталась сама: 3% = $360")], size=11.5, gap=Inches(0.68))
rows = [["Кто", "Что видит"],
        [("Нурлан (партнёр)", INK, True), "Только своих рефералов, их статусы, свою комиссию"],
        [("Менеджер", INK, True), "Обычный лид с источником «партнёр: Нурлан»"],
        [("РОП", INK, True), "Канал «партнёры» в ROI: лиды, конверсия, выручка"],
        [("Админ", INK, True), "Сумму комиссий к выплате за период"]]
make_table(s, Inches(6.8), Inches(1.6), Inches(5.8), rows, [0.8, 1.7], size=10.5, row_h=Pt(26))
card(s, Inches(6.8), Inches(4.4), Inches(5.8), Inches(1.85))
txt(s, Inches(7.05), Inches(4.55), Inches(5.3), Inches(1.6),
    [("ЗАЧЕМ ЭТО БИЗНЕСУ", 10, OKL, True),
     ("Партнёрская сеть = бесплатный канал лидов с оплатой за результат.", 11.5, INK2, False),
     ("Преподаватели, выпускники, блогеры — все через одну систему, без чатов и табличек.", 11.5, INK2, False)])
footer(s, "IX · История 3")

# ============================================================
# 33 · ACTION MAP
# ============================================================
s = add_slide()
kicker(s, "X · Карта действий")
title(s, "«Хочу сделать X — куда идти?»")
rows = [["Хочу…", "Куда идти", "Кто может"],
        ["Изменить цены/тексты на сайте", "/admin → Контент сайта", "Админ"],
        ["Добавить вуз или страну", "/admin → Страны и Университеты", "Админ"],
        ["Создать сотрудника", "/admin → Менеджеры", "Админ"],
        ["Перестроить воронку", "/admin → Статусы лидов", "Админ"],
        ["Шаблон WhatsApp", "/admin → Шаблоны ответов", "Админ"],
        ["«Японские лиды — Анне»", "/admin → Авто-распределение", "Админ"],
        ["Рейтинг команды", "/admin → Leaderboard", "Админ · РОП"],
        ["Подготовиться к 1-на-1", "/admin → 1-на-1 отчёт", "РОП"],
        ["Обработать лида", "/lidy → карточка лида", "Менеджер · РОП"],
        ["Сменить статус", "Карточка → Обзор, или канбан", "Менеджер"],
        ["Передать лид коллеге", "Карточка → Обзор → Передача", "Менеджер"],
        ["Загрузить паспорт клиента", "Карточка → 📎 Файлы", "Менеджер"],
        ["Выгрузить лиды в Excel", "/admin → Дашборд → Экспорт", "Админ · РОП"],
        ["Скачать резервную копию", "/admin → Утилиты и бэкапы", "Админ"]]
make_table(s, Inches(0.7), Inches(1.5), Inches(11.9), rows, [1.3, 1.3, 0.7], size=10.5, row_h=Pt(22))
footer(s, "X · Карта действий")

# ============================================================
# 34 · CHECKLISTS
# ============================================================
s = add_slide()
kicker(s, "XI · Чек-листы")
title(s, "Ритуалы по ролям")
checks = [("Менеджер · день",
           ["09:00 открыть CRM", "Включить «🟢 в сети»", "Глянуть «Мой день»", "Включить Inbox 0",
            "Закрыть просроченные задачи", "Новые лиды — ответ в 15 минут", "Шаблоны вместо ручного набора",
            "Суммы сделок — сразу", "Обед → «не в сети» → обратно", "18:00 Inbox 0 пуст → «не в сети»"]),
          ("РОП · неделя",
           ["Ежедневно 09:00 Telegram-отчёт", "Светофор в админке", "Пн: Leaderboard за неделю",
            "Пн: раздать горячих, топ-3", "Ср: 1-на-1 по готовым отчётам", "Пт: когорты + ROI каналов",
            "Пт: проверить зависшие лиды", "По необходимости: авто-сценарии"]),
          ("Админ · месяц",
           ["1-е число: бэкап в облако", "1-е число: место в БД", "15-е: журнал аудита",
            "Квартал: сменить пароль админки", "Полгода: ревизия правил", "По запросу: контент сайта",
            "По росту: новые сотрудники"])]
for i, (name, items) in enumerate(checks):
    x = Inches(0.7) + Inches(4.05) * i
    card(s, x, Inches(1.55), Inches(3.85), Inches(3.9))
    txt(s, x + Inches(0.22), Inches(1.7), Inches(3.5), Inches(0.4), [(name, 14, INK, True)])
    bullet_list(s, x + Inches(0.22), Inches(2.2), Inches(3.5), Inches(3.2), items, size=10, gap_para=False)
card(s, Inches(0.7), Inches(5.7), Inches(11.9), Inches(1.2), fill=RGBColor(0x2A, 0x1E, 0x0A), line_c=RGBColor(0x78, 0x50, 0x0F))
txt(s, Inches(0.95), Inches(5.85), Inches(11.4), Inches(0.95),
    [("ТРИ «НИКОГДА»", 10, WARNL, True),
     ("1. Не удалять менеджеров — только архивировать.  2. Не забывать «Сохранить» после правок контента.  3. Не оставлять «🟢 в сети» уходя домой.", 11.5, INK, False)])
footer(s, "XI · Чек-листы")

# ============================================================
# 35 · FAQ
# ============================================================
s = add_slide()
kicker(s, "XI · Чек-листы")
title(s, "Частые вопросы")
faqs_l = [("Менеджер забыл пароль?", "Админ выдаёт стартовый пароль заново через «Менеджеры»."),
          ("Лид никому не назначен?", "Все офлайн. Включите «🟢 в сети» — очередь раздастся."),
          ("Случайно удалили лид?", "Восстановление из ежедневного бэкапа — к админу."),
          ("Клиент пришёл повторно?", "Дедупликация привяжет заявку к старой карточке.")]
faqs_r = [("Можно работать с телефона?", "Да: CRM ставится как приложение, push туда же."),
          ("Сайт упал — лиды теряются?", "Хостинг с авто-перезапуском; WhatsApp-кнопка дублирует канал."),
          ("Новый источник рекламы?", "«Варианты источников» в админке + правило распределения."),
          ("Двое правят сайт сразу?", "Сработает последнее «Сохранить». Контент меняет один человек.")]
for col, faqs, x in ((0, faqs_l, Inches(0.7)), (1, faqs_r, Inches(6.8))):
    for i, (q, a) in enumerate(faqs):
        y = Inches(1.6) + Inches(1.25) * i
        card(s, x, y, Inches(5.8), Inches(1.1))
        txt(s, x + Inches(0.22), y + Inches(0.1), Inches(5.4), Inches(0.9),
            [(q, 12.5, INK, True), (a, 10.5, INK2, False)])
footer(s, "XI · FAQ")

# ============================================================
# 36 · FINAL
# ============================================================
s = add_slide(bg=RGBColor(0x0C, 0x2D, 0x52))
txt(s, Inches(0.8), Inches(0.7), Inches(8), Inches(0.4), [("ФИНАЛ", 11, SKYL, True)])
txt(s, Inches(0.8), Inches(1.1), Inches(11), Inches(1.0), [("Готово к работе", 44, WHITE, True)])
txt(s, Inches(0.8), Inches(2.15), Inches(11.5), Inches(0.8),
    [("Платформа закрывает весь цикл: от клика по рекламе до студента в кампусе.", 16, INK2, False),
     ("Детали — в полной базе знаний (PDF, 40 стр.).", 16, INK2, False)])
cols = [("Доступы", ["Сайт: goglobal.kg", "Админка: goglobal.kg/admin", "CRM: goglobal.kg/lidy"]),
        ("Документы", ["Эта презентация (HTML + PPTX)", "База знаний (PDF, 40 стр.)", "Мануал CRM (PDF, 97 стр.)", "Спецификация сайта (PDF)"]),
        ("Первые шаги", ["Менеджерам: раздел VI + история 1", "РОПу: разделы VII и XI", "Админу: всё + база знаний"])]
for i, (name, items) in enumerate(cols):
    x = Inches(0.8) + Inches(4.1) * i
    txt(s, x, Inches(3.5), Inches(3.9), Inches(0.4), [(name.upper(), 12, SKYL, True)])
    bullet_list(s, x, Inches(4.0), Inches(3.9), Inches(2.2), items, size=12, color=RGBColor(0xCB, 0xD5, 0xE1), bcolor=SKYL)
txt(s, Inches(0.8), SLIDE_H - Inches(0.55), Inches(11.7), Inches(0.35),
    [("GoGlobal · база знаний · Спасибо!", 10, INK2, False)])

_target = r"C:\Users\ishem\Downloads\go-global---education-abroad\docs\goglobal-presentation.pptx"
try:
    prs.save(_target)
except PermissionError:
    # file is open in PowerPoint — save next to it
    _target = _target.replace(".pptx", "-v5.pptx")
    prs.save(_target)
print(f"OK: {_n[0]} slides -> {_target}")
